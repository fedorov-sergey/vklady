#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка данных для сайта: Вклады.xlsx -> docs/data.json.

Модель данных: ОДНА СТРОКА = сценарий (продукт x срок x профиль клиента).
База + надбавки выражаются отдельными строками (траты 0 / от 20к / от 100к),
подписка (нет / Приват / ...), «новые деньги» у ГПБ (поле new_money).

Правила нормализации:
  * проценты: 12.39 -> 12.39, 0.12 -> 12, 0.095 -> 9.5
  * формулы вида '=(...)' и '=A*N*I/365' -> None (берём следующий источник ставки)
  * срок: < 60 -> месяцы, >= 60 -> дни (ГПБ хранит в днях)
  * ставка выбирается в порядке: при открытии > с капитализацией при открытии
    > с капитализацией > без капитализации > надбавка «новые деньги»

Запуск:  python build.py
Требования: openpyxl
"""
from __future__ import annotations

import io
import json
import os
import sys
from datetime import datetime

if os.name == "nt":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(SCRIPT_DIR, "Вклады.xlsx")
OUT_PATH = os.path.join(SCRIPT_DIR, "docs", "data.json")

HEADERS = {
    "amount": "Планирую положить",
    "date": "Дата",
    "bank": "Банк",
    "vid": "Вид",
    "product": "Название продукта",
    "popolnenie": "Пополнение",
    "snyatie": "Снятие",
    "min_sum": "от суммы",
    "duration": "Продолжительность",
    "rate_plain": "Ставка без Капитализации",
    "rate_cap": "Ставка с капитализацией",
    "depends_cb": "Зависит от ЦБ",
    "delta_cb": "Разница от ЦБ",
    "rate_cap_open": "Ставка с капитализацией при открытии",
    "account_type": "Тип счета",
    "depends_spending": "Зависит от трат",
    "spending": "Тратить",
    "note": "Прочее",
    "rate_open": "Ставка при открытии",
    "subscription": "Подписка",
    "new_base": "Ставка Новые деньги База",
    "new_bonus": "Ставка Новые деньги Надбавка",
    "income_note": "Обещает Комментарий",
}


def to_percent(v):
    """Любое представление процента -> число в процентах. Формулы -> None."""
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.startswith("="):
        return None
    is_pct_str = "%" in s
    s = s.replace("%", "").strip()
    try:
        f = float(s)
    except ValueError:
        return None
    if not is_pct_str and -1 < f <= 1:  # десятичная доля 0.12 -> 12
        f *= 100
    return round(f, 4)


def parse_delta(v):
    """Разница от ключевой ставки ЦБ: '-1.75', '=-0.8%', 0.003 -> проценты."""
    if v is None:
        return None
    s = str(v).strip().lstrip("=").strip()
    is_pct = "%" in s
    s = s.replace("%", "").strip()
    try:
        f = float(s)
    except ValueError:
        return None
    if not is_pct and -1 < f <= 1:
        f *= 100
    return round(f, 4)


def parse_duration(v):
    """-> (months_or_None, days_or_None, display). <60 считается месяцами."""
    if v is None:
        return None, None, None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None, None, None
    if f < 60:
        months = int(f)
        return months, None, f"{months} мес"
    days = int(f)
    months = int(round(days / 30.4375))
    return months, days, f"{days} дн (~{months} мес)"


def to_num(v, default=None):
    if v is None:
        return default
    try:
        f = float(str(v).replace(" ", "").replace(",", "."))
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return default


_RATE_PRIORITY = [
    ("rate_open", "Ставка при открытии"),
    ("rate_cap_open", "Ставка с капитализацией при открытии"),
    ("rate_cap", "Ставка с капитализацией"),
    ("rate_plain", "Ставка без Капитализации"),
    ("new_bonus", "Ставка Новые деньги Надбавка"),
]


def normalize_rows(raw_rows):
    rows = []
    for i, r in enumerate(raw_rows, start=2):
        product = (r.get("product") or "").strip()
        rate = None
        rate_note = None
        for field, label in _RATE_PRIORITY:
            p = to_percent(r.get(field))
            if p is not None:
                rate = p
                rate_note = label
                break
        if rate is None:
            continue  # строка без ставки (заглушка/не заполнена)

        amount = to_num(r.get("amount"), 0)
        min_sum = to_num(r.get("min_sum"), 0)
        months, days, duration_disp = parse_duration(r.get("duration"))

        vid = (r.get("vid") or "").strip()
        vtype = "счет" if vid == "Счет" else "вклад"

        spending = None
        if str(r.get("depends_spending") or "").strip() == "да":
            spending = to_num(r.get("spending"))

        sub_raw = (r.get("subscription") or "").strip()
        subscription = None if sub_raw.lower() in ("", "нет", "-") else sub_raw

        new_money = to_percent(r.get("new_bonus")) is not None
        rate_base = to_percent(r.get("new_base")) if new_money else None

        depends_cb = (r.get("depends_cb") or "").strip().lower() in ("да", "yes", "+")
        cbr = None
        if depends_cb:
            cbr = {
                "delta": parse_delta(r.get("delta_cb")),
                "rate_at_open": to_percent(r.get("rate_cap_open")),
            }

        date_val = r.get("date")
        date_iso = date_val.date().isoformat() if isinstance(date_val, datetime) else None

        rows.append({
            "id": f"R{i}",
            "bank": (r.get("bank") or "").strip() or "?",
            "product": product or "?",
            "type": vtype,
            "date": date_iso,
            "default_amount": amount,
            "min_sum": min_sum,
            "popolnenie": (r.get("popolnenie") or "").strip() or None,
            "snyatie": (r.get("snyatie") or "").strip() or None,
            "duration_months": months,
            "duration_days": days,
            "duration_display": duration_disp,
            "rate": rate,
            "rate_base": rate_base,
            "rate_note": rate_note,
            "cbr": cbr,
            "account_type": (r.get("account_type") or "").strip() or None,
            "profile": {
                "spending": spending,
                "subscription": subscription,
                "new_money": new_money,
            },
            "note": (r.get("note") or "").strip() or None,
            "income_note": (r.get("income_note") or "").strip() or None,
        })
    return rows


def main():
    import openpyxl  # локальный импорт: чтобы ошибки сборки были понятны

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)  # значения, не формулы
    ws = wb[wb.sheetnames[0]]

    header_row = [c.value for c in ws[1]]
    col_index = {name: i for i, name in enumerate(header_row)}
    field_to_col = {field: col_index.get(label) for field, label in HEADERS.items()}

    raw_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        d = {}
        for field, col in field_to_col.items():
            if col is not None and col < len(row):
                d[field] = row[col]
        raw_rows.append(d)

    rows = normalize_rows(raw_rows)

    # оставляем только срез за последнюю дату (защита от накопления истории в файле)
    dates = sorted({r["date"] for r in rows if r["date"]})
    if dates:
        latest = dates[-1]
        rows = [r for r in rows if r.get("date") == latest]
    latest = dates[-1] if dates else None

    banks = sorted({r["bank"] for r in rows})
    products = sorted({f'{r["bank"]} · {r["product"]}' for r in rows})
    durations = sorted({r["duration_months"] for r in rows if r["duration_months"]})
    spendings = sorted({r["profile"]["spending"] for r in rows if r["profile"]["spending"] is not None})
    subscriptions = sorted(
        {r["profile"]["subscription"] for r in rows if r["profile"]["subscription"] is not None}
    )
    min_sum = min((r["min_sum"] for r in rows if r.get("min_sum")), default=0)
    max_sum = max(r["min_sum"] for r in rows)
    default_amount = max({r["default_amount"] for r in rows if r.get("default_amount")}, default=1000000)

    out = {
        "meta": {
            "site": "Выбор вклада — сравнение ставок",
            "source": os.path.basename(XLSX_PATH),
            "updated": latest,
            "generated": datetime.now().strftime("%Y-%m-%d"),
            "default_amount": default_amount,
        },
        "filters": {
            "banks": banks,
            "products": products,
            "durations": durations,
            "spending_options": spendings,
            "subscription_options": subscriptions,
            "min_sum": min_sum,
            "max_sum": max_sum,
        },
        "rows": rows,
    }

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    n = len(rows)
    print(f"OK: {n} сценариев, банков: {len(banks)}, продуктов: {len(products)}")
    print(f"Дата среза: {latest} | сроки (мес): {durations} | траты: {spendings}")
    print(f"Ставки: {min(r['rate'] for r in rows)}..{max(r['rate'] for r in rows)}%")
    print(f"-> {OUT_PATH}")


if __name__ == "__main__":
    main()